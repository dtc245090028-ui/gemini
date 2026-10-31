"""API Endpoints cho tra cứu Thẻ kho (Stock Ledger), Xuất Excel và Điều chỉnh kiểm kê tồn kho (Stock Adjustment)."""

import calendar
from datetime import date, datetime, timedelta, timezone
import io
from typing import List, Optional, Union
import time
from fastapi import APIRouter, Depends, Header, HTTPException, Query, Request, status
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from sqlalchemy import case, func, or_
from sqlalchemy.orm import Session, joinedload

from app.api.deps import get_current_user, get_db, require_roles
from app.models.export_note import ExportNote, ExportNoteDetail
from app.models.import_note import ImportNote, ImportNoteDetail
from app.models.product import Product
from app.models.stock_ledger import StockLedger
from app.models.user import User
from app.schemas.stock_ledger import (
    ReferenceDetailResponse,
    ReferenceLedgerItemResponse,
    StockAdjustmentCreate,
    StockAdjustmentResponse,
    StockLedgerListResponse,
    StockLedgerResponse,
    StockLedgerSummary,
)
from app.services.inventory_service import adjust_stock
from app.core.write_gate import WriteGateContext
from app.core.config import VN_TZ, get_now_vn

router = APIRouter(prefix="/stock-ledger", tags=["Thẻ kho & Điều chỉnh (Stock Ledger)"])


def _resolve_date_bounds(
    from_date: Optional[datetime],
    to_date: Optional[datetime],
    date_from: Optional[date],
    date_to: Optional[date],
) -> tuple[Optional[datetime], Optional[datetime]]:
    """Chuẩn hóa khoảng lọc thời gian. date_to bao gồm trọn ngày cuối (< date_to + 1 ngày)."""
    start_dt: Optional[datetime] = None
    end_dt_exclusive: Optional[datetime] = None

    if date_from is not None:
        start_dt = datetime(date_from.year, date_from.month, date_from.day, 0, 0, 0)
    elif from_date is not None:
        start_dt = (
            from_date.astimezone(VN_TZ).replace(tzinfo=None)
            if from_date.tzinfo
            else from_date
        )

    if date_to is not None:
        end_dt_exclusive = datetime(date_to.year, date_to.month, date_to.day, 0, 0, 0) + timedelta(days=1)
    elif to_date is not None:
        raw_to = (
            to_date.astimezone(VN_TZ).replace(tzinfo=None)
            if to_date.tzinfo
            else to_date
        )
        end_dt_exclusive = raw_to + timedelta(microseconds=1)

    return start_dt, end_dt_exclusive


@router.get("", response_model=Union[List[StockLedgerResponse], StockLedgerListResponse], include_in_schema=False)
@router.get(
    "/",
    response_model=Union[List[StockLedgerResponse], StockLedgerListResponse],
    summary="Tra cứu lịch sử Thẻ kho (Audit Trail)",
)
def get_stock_ledger_entries(
    product_id: Optional[int] = Query(None, description="Lọc theo ID mặt hàng"),
    transaction_type: Optional[str] = Query(None, description="Lọc theo loại giao dịch: IMPORT, EXPORT, ADJUSTMENT"),
    q: Optional[str] = Query(None, description="Tìm gần đúng theo mã chứng từ, ghi chú, tên mặt hàng"),
    date_from: Optional[date] = Query(None, description="Từ ngày (YYYY-MM-DD)"),
    date_to: Optional[date] = Query(None, description="Đến ngày (YYYY-MM-DD, bao gồm cả ngày cuối)"),
    from_date: Optional[datetime] = Query(None, description="Từ ngày (ISO datetime)"),
    to_date: Optional[datetime] = Query(None, description="Đến ngày (ISO datetime)"),
    include_summary: bool = Query(False, description="Nếu True, trả về {items, total, summary}"),
    skip: int = Query(0, ge=0, description="Số lượng bỏ qua"),
    limit: int = Query(100, ge=1, le=1000, description="Số lượng lấy tối đa"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Truy vấn sổ kiểm toán Thẻ kho phục vụ truy vết lịch sử biến động từng giây."""
    start_dt, end_dt_exclusive = _resolve_date_bounds(from_date, to_date, date_from, date_to)

    base_query = db.query(StockLedger).join(Product, StockLedger.product_id == Product.id)

    if product_id is not None:
        base_query = base_query.filter(StockLedger.product_id == product_id)
    if transaction_type:
        base_query = base_query.filter(StockLedger.transaction_type == transaction_type.strip().upper())
    if start_dt is not None:
        base_query = base_query.filter(StockLedger.transaction_date >= start_dt)
    if end_dt_exclusive is not None:
        base_query = base_query.filter(StockLedger.transaction_date < end_dt_exclusive)
    if q and q.strip():
        kw = f"%{q.strip()}%"
        base_query = base_query.filter(
            or_(
                StockLedger.reference_code.ilike(kw),
                StockLedger.note.ilike(kw),
                Product.name.ilike(kw),
                Product.code.ilike(kw),
            )
        )

    total_count = base_query.count()

    # Tính khối summary khi lọc đúng 1 mặt hàng (product_id có giá trị)
    summary_obj: Optional[StockLedgerSummary] = None
    if product_id is not None:
        agg_row = (
            base_query.with_entities(
                func.coalesce(
                    func.sum(
                        case(
                            (StockLedger.quantity_change > 0, StockLedger.quantity_change),
                            else_=0,
                        )
                    ),
                    0,
                ).label("total_in"),
                func.coalesce(
                    func.sum(
                        case(
                            (StockLedger.quantity_change < 0, -StockLedger.quantity_change),
                            else_=0,
                        )
                    ),
                    0,
                ).label("total_out"),
            )
            .first()
        )
        total_in = int(agg_row.total_in if agg_row else 0)
        total_out = int(agg_row.total_out if agg_row else 0)

        if start_dt is not None:
            # Có date_from: lấy balance_after của dòng gần nhất TRƯỚC start_dt
            prev_entry = (
                db.query(StockLedger)
                .filter(
                    StockLedger.product_id == product_id,
                    StockLedger.transaction_date < start_dt,
                )
                .order_by(StockLedger.transaction_date.desc(), StockLedger.id.desc())
                .first()
            )
            opening_balance = int(prev_entry.balance_after) if prev_entry else 0
        else:
            # Không có date_from: lấy balance_before của dòng cũ nhất thỏa bộ lọc
            oldest_entry = (
                base_query.order_by(StockLedger.transaction_date.asc(), StockLedger.id.asc())
                .first()
            )
            if oldest_entry:
                opening_balance = int(oldest_entry.balance_after - oldest_entry.quantity_change)
            else:
                opening_balance = 0

        closing_balance = opening_balance + total_in - total_out
        summary_obj = StockLedgerSummary(
            opening_balance=opening_balance,
            total_in=total_in,
            total_out=total_out,
            closing_balance=closing_balance,
        )

    entries = (
        base_query.options(
            joinedload(StockLedger.product),
            joinedload(StockLedger.creator),
        )
        .order_by(StockLedger.transaction_date.desc(), StockLedger.id.desc())
        .offset(skip)
        .limit(limit)
        .all()
    )

    results: List[StockLedgerResponse] = []
    for e in entries:
        qty_change = int(e.quantity_change)
        bal_after = int(e.balance_after)
        qty_in = max(qty_change, 0)
        qty_out = max(-qty_change, 0)
        bal_before = bal_after - qty_change
        prod_min_stock = int(e.product.min_stock) if e.product and e.product.min_stock is not None else 0
        prod_code = e.product.code if e.product else None

        results.append(
            StockLedgerResponse(
                id=e.id,
                product_id=e.product_id,
                transaction_type=e.transaction_type,
                reference_code=e.reference_code,
                quantity_change=qty_change,
                balance_after=bal_after,
                created_by=e.created_by,
                transaction_date=e.transaction_date,
                note=e.note,
                product_name=e.product.name if e.product else None,
                product_code=prod_code,
                product_sku=prod_code,
                creator_name=e.creator.full_name if e.creator else None,
                qty_in=qty_in,
                qty_out=qty_out,
                balance_before=bal_before,
                min_stock=prod_min_stock,
                summary=summary_obj,
            )
        )

    if include_summary:
        return StockLedgerListResponse(
            items=results,
            total=total_count,
            summary=summary_obj,
        )

    return results


@router.get(
    "/by-reference/{reference_code}",
    response_model=ReferenceDetailResponse,
    summary="Lấy chi tiết chứng từ và toàn bộ các dòng Thẻ kho cùng mã chứng từ",
)
def get_stock_ledger_by_reference(
    reference_code: str,
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Truy vấn toàn bộ các mặt hàng biến động trong cùng một mã chứng từ (PN-..., PX-..., ADJ-..., CANCEL-...)."""
    ref_clean = reference_code.strip()
    entries = (
        db.query(StockLedger)
        .options(
            joinedload(StockLedger.product),
            joinedload(StockLedger.creator),
        )
        .filter(StockLedger.reference_code == ref_clean)
        .order_by(StockLedger.id.asc())
        .all()
    )

    if not entries:
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND,
            detail=f"Không tìm thấy dữ liệu thẻ kho cho mã chứng từ '{ref_clean}'.",
        )

    first_entry = entries[0]
    partner_name: Optional[str] = None
    note_status: Optional[str] = None
    total_amount: Optional[float] = None
    note_text: Optional[str] = first_entry.note
    price_map: dict[int, tuple[float, float]] = {}

    # Tìm phiếu gốc (ImportNote hoặc ExportNote) nếu có
    lookup_code = ref_clean[7:] if ref_clean.startswith("CANCEL-") else ref_clean

    imp_note = (
        db.query(ImportNote)
        .options(joinedload(ImportNote.supplier), joinedload(ImportNote.details))
        .filter(ImportNote.code == lookup_code)
        .first()
    )
    if imp_note:
        partner_name = imp_note.supplier.name if imp_note.supplier else None
        note_status = imp_note.status
        total_amount = float(imp_note.total_amount) if imp_note.total_amount is not None else None
        if imp_note.note and not ref_clean.startswith("CANCEL-"):
            note_text = imp_note.note
        for d in imp_note.details:
            price_map[d.product_id] = (float(d.unit_price), float(d.subtotal))
    else:
        exp_note = (
            db.query(ExportNote)
            .options(joinedload(ExportNote.details))
            .filter(ExportNote.code == lookup_code)
            .first()
        )
        if exp_note:
            partner_name = exp_note.recipient_name
            note_status = exp_note.status
            total_amount = float(exp_note.total_amount) if exp_note.total_amount is not None else None
            if exp_note.note and not ref_clean.startswith("CANCEL-"):
                note_text = exp_note.note
            for d in exp_note.details:
                price_map[d.product_id] = (float(d.unit_price), float(d.subtotal))

    # Nếu không tìm thấy đối tác từ phiếu gốc thì trích xuất từ ghi chú nếu có
    if not partner_name and first_entry.note:
        for prefix in ("Nhập kho từ ", "Xuất kho cho ", "Xuất kho giao cho "):
            if first_entry.note.startswith(prefix):
                extracted = first_entry.note[len(prefix):].split(" theo phiếu ")[0].strip()
                if extracted:
                    partner_name = extracted
                break

    items: List[ReferenceLedgerItemResponse] = []
    total_qty_in = 0
    total_qty_out = 0

    for e in entries:
        qty_change = int(e.quantity_change)
        bal_after = int(e.balance_after)
        q_in = max(qty_change, 0)
        q_out = max(-qty_change, 0)
        bal_before = bal_after - qty_change
        total_qty_in += q_in
        total_qty_out += q_out

        u_price: Optional[float] = None
        sub_val: Optional[float] = None
        if e.product_id in price_map:
            u_price, _ = price_map[e.product_id]
            sub_val = float(abs(qty_change) * u_price)

        prod_code = e.product.code if e.product else None
        items.append(
            ReferenceLedgerItemResponse(
                id=e.id,
                product_id=e.product_id,
                product_code=prod_code,
                product_sku=prod_code,
                product_name=e.product.name if e.product else f"Mặt hàng #{e.product_id}",
                unit=e.product.unit if e.product else None,
                transaction_type=e.transaction_type,
                quantity_change=qty_change,
                qty_in=q_in,
                qty_out=q_out,
                balance_before=bal_before,
                balance_after=bal_after,
                min_stock=int(e.product.min_stock) if e.product and e.product.min_stock is not None else 0,
                unit_price=u_price,
                subtotal=sub_val,
            )
        )

    return ReferenceDetailResponse(
        reference_code=ref_clean,
        transaction_type=first_entry.transaction_type,
        transaction_date=first_entry.transaction_date,
        created_by=first_entry.created_by,
        creator_name=first_entry.creator.full_name if first_entry.creator else "Hệ thống",
        partner_name=partner_name,
        note=note_text,
        note_status=note_status,
        total_amount=total_amount,
        total_qty_in=total_qty_in,
        total_qty_out=total_qty_out,
        items=items,
    )


@router.get(
    "/export-excel",
    summary="Xuất sổ cái Thẻ kho ra file Excel (.xlsx) gồm 2 sheet Nhập kho và Xuất kho",
)
def export_stock_ledger_excel(
    period: str = Query(
        "day",
        description="Phạm vi thời gian: day (Theo ngày), week (Theo tuần), month (Theo tháng)",
    ),
    target_date: Optional[date] = Query(
        None,
        description="Ngày mốc định dạng YYYY-MM-DD (Mặc định: ngày hôm nay theo giờ Việt Nam)",
    ),
    product_id: Optional[int] = Query(None, description="Lọc theo ID mặt hàng (nếu có)"),
    transaction_type: Optional[str] = Query(None, description="Lọc theo loại giao dịch"),
    q: Optional[str] = Query(None, description="Từ khóa tìm kiếm"),
    date_from: Optional[date] = Query(None, description="Từ ngày tùy chọn"),
    date_to: Optional[date] = Query(None, description="Đến ngày tùy chọn"),
    split_columns: bool = Query(False, description="Xuất bảng với 2 cột Nhập / Xuất tách biệt"),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
):
    """Xuất dữ liệu Thẻ kho trong phạm vi Ngày / Tuần / Tháng ra file .xlsx."""

    now_vn = datetime.now(VN_TZ)
    chosen_date = target_date or now_vn.date()
    mode = (period or "day").strip().lower()

    if date_from is not None or date_to is not None:
        start_day = date_from or date(2020, 1, 1)
        end_day = date_to or now_vn.date()
        period_label = f"Từ {start_day.strftime('%d/%m/%Y')} đến {end_day.strftime('%d/%m/%Y')}"
        file_suffix = f"tu_{start_day.strftime('%Y%m%d')}_den_{end_day.strftime('%Y%m%d')}"
    elif mode == "day":
        start_day = chosen_date
        end_day = chosen_date
        period_label = f"Ngày {chosen_date.strftime('%d/%m/%Y')}"
        file_suffix = f"ngay_{chosen_date.strftime('%Y%m%d')}"
    elif mode == "week":
        start_day = chosen_date - timedelta(days=chosen_date.weekday())
        end_day = start_day + timedelta(days=6)
        period_label = f"Tuần {start_day.strftime('%d/%m/%Y')} - {end_day.strftime('%d/%m/%Y')}"
        file_suffix = f"tuan_{start_day.strftime('%Y%m%d')}_{end_day.strftime('%Y%m%d')}"
    elif mode == "month":
        start_day = date(chosen_date.year, chosen_date.month, 1)
        last_day_num = calendar.monthrange(chosen_date.year, chosen_date.month)[1]
        end_day = date(chosen_date.year, chosen_date.month, last_day_num)
        period_label = f"Tháng {chosen_date.month:02d}/{chosen_date.year}"
        file_suffix = f"thang_{chosen_date.year}{chosen_date.month:02d}"
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Tham số 'period' không hợp lệ. Chỉ chấp nhận: day, week, month.",
        )

    start_dt = datetime(start_day.year, start_day.month, start_day.day, 0, 0, 0)
    end_dt_exclusive = datetime(end_day.year, end_day.month, end_day.day, 0, 0, 0) + timedelta(days=1)

    query = (
        db.query(StockLedger)
        .join(Product, StockLedger.product_id == Product.id)
        .options(
            joinedload(StockLedger.product),
            joinedload(StockLedger.creator),
        )
        .filter(
            StockLedger.transaction_date >= start_dt,
            StockLedger.transaction_date < end_dt_exclusive,
        )
    )
    if product_id is not None:
        query = query.filter(StockLedger.product_id == product_id)
    if transaction_type:
        query = query.filter(StockLedger.transaction_type == transaction_type.strip().upper())
    if q and q.strip():
        kw = f"%{q.strip()}%"
        query = query.filter(
            or_(
                StockLedger.reference_code.ilike(kw),
                StockLedger.note.ilike(kw),
                Product.name.ilike(kw),
                Product.code.ilike(kw),
            )
        )

    entries = query.order_by(StockLedger.transaction_date.desc(), StockLedger.id.desc()).all()

    import_entries = [
        e for e in entries
        if e.transaction_type == "IMPORT" or (e.transaction_type == "ADJUSTMENT" and e.quantity_change > 0)
    ]
    export_entries = [
        e for e in entries
        if e.transaction_type == "EXPORT" or (e.transaction_type == "ADJUSTMENT" and e.quantity_change < 0)
    ]

    wb = Workbook()
    ws_import = wb.active
    ws_import.title = "Nhập kho"
    ws_export = wb.create_sheet(title="Xuất kho")

    title_font = Font(name="Arial", size=13, bold=True, color="1E293B")
    subtitle_font = Font(name="Arial", size=10, italic=True, color="475569")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    cell_font = Font(name="Arial", size=10, color="0F172A")
    bold_cell_font = Font(name="Arial", size=10, bold=True, color="0F172A")

    thin_side = Side(border_style="thin", color="CBD5E1")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    if split_columns:
        headers = [
            "Thời gian",
            "Tên hàng hóa",
            "Loại nghiệp vụ",
            "Mã chứng từ",
            "Nhập",
            "Xuất",
            "Tồn sau giao dịch",
            "Người thực hiện",
            "Ghi chú",
        ]
        col_widths = [22, 38, 18, 22, 14, 14, 18, 24, 42]
        last_col_letter = "I"
    else:
        headers = [
            "Thời gian",
            "Tên hàng hóa",
            "Mã chứng từ",
            "Số lượng",
            "Tồn sau giao dịch",
            "Người thực hiện",
            "Ghi chú",
        ]
        col_widths = [22, 38, 22, 15, 18, 24, 42]
        last_col_letter = "G"

    def populate_sheet(ws, sheet_title: str, rows: List[StockLedger], header_hex: str):
        ws.merge_cells(f"A1:{last_col_letter}1")
        ws["A1"] = f"SỔ CÁI THẺ KHO - {sheet_title.upper()} ({period_label.upper()})"
        ws["A1"].font = title_font
        ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[1].height = 24

        ws.merge_cells(f"A2:{last_col_letter}2")
        ws["A2"] = (
            f"Người xuất báo cáo: {current_user.full_name} ({current_user.username}) | "
            f"Thời điểm xuất: {now_vn.strftime('%d/%m/%Y %H:%M:%S')}"
        )
        ws["A2"].font = subtitle_font
        ws["A2"].alignment = Alignment(horizontal="left", vertical="center")
        ws.row_dimensions[2].height = 18

        header_fill = PatternFill(start_color=header_hex, end_color=header_hex, fill_type="solid")
        ws.row_dimensions[4].height = 24
        for col_idx, h_text in enumerate(headers, start=1):
            cell = ws.cell(row=4, column=col_idx, value=h_text)
            cell.font = header_font
            cell.fill = header_fill
            cell.border = cell_border
            cell.alignment = Alignment(
                horizontal="left" if h_text in ("Tên hàng hóa", "Người thực hiện", "Ghi chú") else "center",
                vertical="center",
            )

        for row_idx, item in enumerate(rows, start=5):
            ws.row_dimensions[row_idx].height = 20
            time_str = item.transaction_date.strftime("%d/%m/%Y %H:%M:%S") if item.transaction_date else "-"
            prod_name = item.product.name if item.product else f"Mặt hàng #{item.product_id}"
            creator_name = item.creator.full_name if item.creator else "Hệ thống"
            qty_change = int(item.quantity_change)
            qty_in = max(qty_change, 0)
            qty_out = max(-qty_change, 0)
            note_text = item.note or "-"

            if split_columns:
                type_label = (
                    "Nhập kho"
                    if item.transaction_type == "IMPORT"
                    else "Xuất kho"
                    if item.transaction_type == "EXPORT"
                    else "Kiểm kê"
                )
                row_values = [
                    time_str,
                    prod_name,
                    type_label,
                    item.reference_code,
                    qty_in if qty_in > 0 else "",
                    qty_out if qty_out > 0 else "",
                    int(item.balance_after),
                    creator_name,
                    note_text,
                ]
                for col_idx, val in enumerate(row_values, start=1):
                    c = ws.cell(row=row_idx, column=col_idx, value=val)
                    c.font = bold_cell_font if col_idx in (5, 6, 7) else cell_font
                    c.border = cell_border
                    if col_idx in (1, 3, 4):
                        c.alignment = Alignment(horizontal="center", vertical="center")
                    elif col_idx in (5, 6, 7):
                        c.alignment = Alignment(horizontal="right", vertical="center")
                        if isinstance(val, int):
                            c.number_format = "#,##0"
                    else:
                        c.alignment = Alignment(horizontal="left", vertical="center")
            else:
                qty = abs(qty_change)
                row_values = [
                    time_str,
                    prod_name,
                    item.reference_code,
                    qty,
                    int(item.balance_after),
                    creator_name,
                    note_text,
                ]
                for col_idx, val in enumerate(row_values, start=1):
                    c = ws.cell(row=row_idx, column=col_idx, value=val)
                    c.font = bold_cell_font if col_idx in (4, 5) else cell_font
                    c.border = cell_border
                    if col_idx in (1, 3):
                        c.alignment = Alignment(horizontal="center", vertical="center")
                    elif col_idx in (4, 5):
                        c.alignment = Alignment(horizontal="right", vertical="center")
                        c.number_format = "#,##0"
                    else:
                        c.alignment = Alignment(horizontal="left", vertical="center")

        for idx, width in enumerate(col_widths, start=1):
            col_letter = ws.cell(row=4, column=idx).column_letter
            ws.column_dimensions[col_letter].width = width

    populate_sheet(ws_import, "Bảng Nhập Kho", import_entries, "166534")
    populate_sheet(ws_export, "Bảng Xuất Kho", export_entries, "1E40AF")

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)

    filename = f"the_kho_{file_suffix}.xlsx"
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": f'attachment; filename="{filename}"',
            "Access-Control-Expose-Headers": "Content-Disposition",
        },
    )


@router.post(
    "/adjust",
    response_model=StockAdjustmentResponse,
    status_code=status.HTTP_200_OK,
    summary="Lập phiếu điều chỉnh kiểm kê tồn kho (Lối thoát khi hủy phiếu bị chặn hoặc kiểm kê thực tế)",
)
def create_stock_adjustment(
    request: Request,
    adj_in: StockAdjustmentCreate,
    x_idempotency_key: Optional[str] = Header(None, alias="X-Idempotency-Key", description="Khóa chống ghi trùng lặp (Idempotency Key)"),
    x_client_id: Optional[str] = Header(None, alias="X-Client-Id", description="Định danh phiên client"),
    db: Session = Depends(get_db),
    current_user: User = Depends(require_roles(["ADMIN", "WAREHOUSE_KEEPER"])),
):
    """Điều chỉnh số lượng tồn kho theo số lượng thực tế sau kiểm kê (Hỗ trợ Idempotency Key & Write Gate)."""
    arrival_t = getattr(request.state, "arrival_time", None)
    st = getattr(request.state, "stage_timing", None)
    if st:
        st.t_handler_enter = time.perf_counter()
    with WriteGateContext(arrival_time=arrival_t, db=db):
        resp = adjust_stock(
            db=db,
            adj_in=adj_in,
            current_user_id=current_user.id,
            idempotency_key=x_idempotency_key,
            stage_timing=st,
            client_id=x_client_id,
        )
    if st:
        st.t_handler_exit = time.perf_counter()
    return resp
